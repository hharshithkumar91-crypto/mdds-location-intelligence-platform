"""
Phase 2 PowerPoint Presentation Generator (16:9 Widescreen)
Generates Phase_2_API_and_Dashboard_Specification.pptx
"""

import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

COLOR_BG = RGBColor(15, 23, 42)          # #0F172A
COLOR_CARD = RGBColor(30, 41, 59)        # #1E293B
COLOR_CARD_BORDER = RGBColor(51, 65, 85) # #334155
COLOR_PRIMARY = RGBColor(37, 99, 235)    # #2563EB
COLOR_ACCENT = RGBColor(96, 165, 250)    # #60A5FA
COLOR_TEXT_WHITE = RGBColor(248, 250, 252)
COLOR_TEXT_MUTED = RGBColor(148, 163, 184)
COLOR_TABLE_HDR = RGBColor(30, 58, 138)

def set_slide_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG

def add_header(slide, title_text, category_text="PHASE 2 API & PORTAL SPECIFICATION"):
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
    tf_cat = cat_box.text_frame
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_ACCENT

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
    tf_title = title_box.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TEXT_WHITE

def add_card(slide, left, top, width, height, title, body_bullets, accent_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_CARD
    shape.line.color.rgb = accent_color or COLOR_CARD_BORDER
    shape.line.width = Pt(1.5)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.25)

    p_title = tf.paragraphs[0]
    p_title.text = title
    p_title.font.size = Pt(16)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_ACCENT
    p_title.space_after = Pt(10)

    for bullet in body_bullets:
        p = tf.add_paragraph()
        p.text = "• " + bullet
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_after = Pt(6)

def build_phase2_presentation(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Slide 1: Cover
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    banner = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.4), Inches(7.5))
    banner.fill.solid()
    banner.fill.fore_color.rgb = COLOR_PRIMARY

    t_box = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(2.5))
    tf = t_box.text_frame
    p0 = tf.paragraphs[0]
    p0.text = "BOLD ANALYTICS CAPSTONE 2026"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT

    p1 = tf.add_paragraph()
    p1.text = "Phase 2: API Development, Admin Dashboard & B2B Portal"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_WHITE

    p2 = tf.add_paragraph()
    p2.text = "REST Endpoints, Standard Response Envelopes, Multi-Tenant Security & Demo Client"
    p2.font.size = Pt(17)
    p2.font.color.rgb = COLOR_TEXT_MUTED

    card_meta = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(4.7), Inches(11.0), Inches(1.8))
    card_meta.fill.solid()
    card_meta.fill.fore_color.rgb = COLOR_CARD
    card_meta.line.color.rgb = COLOR_CARD_BORDER
    tf_m = card_meta.text_frame
    pm1 = tf_m.paragraphs[0]
    pm1.text = "ENGINEERING CONTRIBUTORS & MILESTONE"
    pm1.font.size = Pt(12)
    pm1.font.bold = True
    pm1.font.color.rgb = COLOR_ACCENT
    pm2 = tf_m.add_paragraph()
    pm2.text = "• Team: Harshith Kumar (HH), Riyan Pathan (RI), Geetha (YG)"
    pm2.font.size = Pt(13)
    pm2.font.color.rgb = COLOR_TEXT_WHITE
    pm3 = tf_m.add_paragraph()
    pm3.text = "• Milestone: Phase 2 (Due: Sep 11, 2026) | Verification: 100% Automated API Tests Passed"
    pm3.font.size = Pt(13)
    pm3.font.color.rgb = COLOR_TEXT_MUTED

    # Slide 2: API Specifications (Section 6)
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "API Gateway Specifications (Section 6)", "CORE REST CONTRACTS")
    add_card(s2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8),
             "Standard Response Envelope (Section 6.3)",
             ["Every response returns uniform JSON structure.",
              "success: boolean indicator of transaction state.",
              "count: total records delivered in payload.",
              "meta.requestId: unique tracing UUID (e.g. req_ad7b5b3f).",
              "meta.responseTime: execution latency in milliseconds.",
              "meta.rateLimit: remaining and limit quota tracking."])
    add_card(s2, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8),
             "Core Endpoints & Dropdown Schema (Section 6.4 & 6.5)",
             ["GET /v1/states: Lists all states and union territories.",
              "GET /v1/states/{id}/districts: Districts by state.",
              "GET /v1/districts/{id}/subdistricts: Sub-districts by district.",
              "GET /v1/subdistricts/{id}/villages: Paginated revenue villages.",
              "GET /v1/autocomplete: Formats {value, label, fullAddress, hierarchy} for instant dropdown population."])

    # Slide 3: Admin & B2B Dashboard (Sections 7 & 8)
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Frontend Dashboard & Admin Panel (Sections 7 & 8)", "REACT & RECHARTS ARCHITECTURE")
    add_card(s3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8),
             "Frontend Tech Stack (Section 7)",
             ["React 18+ with TypeScript for strict type checking.",
              "Vite for ultra-fast HMR and optimized production bundles.",
              "Tailwind CSS for responsive design system styling.",
              "Recharts for responsive, declarative SVG visualizations.",
              "Zustand (Global State) + React Query (Server Caching).",
              "Performance SLA: Initial load <2s, charts render <500ms."])
    add_card(s3, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8),
             "Admin Panel Features (Section 8)",
             ["Analytics Visualizations: Bar chart (top states), Line chart (30-day traffic), Area chart (p95/p99 latency).",
              "User Management: Filter by status (Pending, Active) and plan type (Free, Premium, Pro).",
              "User Approval Workflow: Review corporate email and unlock API key generation upon approval.",
              "State Access Matrix: Grant full India, specific states, or regional bundles.",
              "Village Master Browser: Paginated table (500 to 10,000 rows) with CSV export."])

    # Slide 4: Authentication & Rate Limiting (Sections 10 & 11)
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Multi-Tenant Security & Rate Limiting (Sections 10 & 11)", "ENTERPRISE IDENTITY & QUOTAS")
    add_card(s4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8),
             "Authentication & Key Format (Section 10)",
             ["Dual-layer auth: JWT (24h) for web portal sessions; X-API-Key + X-API-Secret for programmatic access.",
              "API Key syntax: ak_[32 hex characters] (e.g. ak_a1b2c3d4...).",
              "API Secret syntax: as_[32 hex characters] (hashed with bcrypt).",
              "Hardened Security Headers: nosniff, DENY, HSTS, CSP enforced on all responses.",
              "Maximum 5 active keys per organization with instant revocation."])
    add_card(s4, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8),
             "Tier Quotas & Rate Limits (Section 11)",
             ["Free Tier ($0): 5,000 requests/day, 100 req/min burst.",
              "Premium Tier ($49): 50,000 requests/day, 500 req/min burst.",
              "Pro Tier ($199): 300,000 requests/day, 2,000 req/min burst.",
              "Unlimited Tier ($499): 1,000,000 requests/day, 5,000 req/min burst.",
              "Standard Headers: X-RateLimit-Limit, Remaining, Reset.",
              "Automated Alerts: Email notifications at 80% and 95% of daily limit."])

    # Slide 5: Demo Client Project (Section 13)
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Demo Client Project (Section 13)", "B2B REFERENCE IMPLEMENTATION")
    add_card(s5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8),
             "Reference Web App Architecture (demo-client/)",
             ["Simple contact form application showcasing live API integration.",
              "Inputs: Full Name, Work Email, Phone, Village / Locality, Message.",
              "Live Autocomplete: Queries /v1/autocomplete on >= 2 characters.",
              "Auto-Fill Mechanics: Selecting a village automatically populates Sub-District, District, State, and Country ('India').",
              "Standardized Submission: Form submit payload contains verified standardized address."])
    add_card(s5, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8),
             "Demo Configuration & Security Guardrails",
             ["Public Demo Key: demo_public_key_for_presentations.",
              "Strict Read-Only Access: Write operations rejected with 403 ACCESS_DENIED.",
              "Restricted Quota: Capped at 100 requests per day to prevent abuse.",
              "Regional Scoping: Restricted strictly to Maharashtra state data.",
              "Zero Auth Required for Presentations: Ready for client sales demos."])

    # Slide 6: Verification Results
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Automated Verification Results", "100% SPECIFICATION COMPLIANCE")
    add_card(s6, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8),
             "Automated Test Suite Output (python scripts/test_phase2_apis.py)",
             ["[PASS] Server online & healthy. HTTP 200",
              "[PASS] Standard Response Envelope Validated (requestId, responseTime: 1ms, rateLimit: 99/100)",
              "[PASS] GET /v1/states/st_27/districts -> Found Nandurbar (code: 497)",
              "[PASS] GET /v1/districts/dt_497/subdistricts -> Found Akkalkuwa (preserved code: '03950')",
              "[PASS] Section 6.5 Autocomplete Schema -> { value, label: 'Manibeli', fullAddress, hierarchy }",
              "[PASS] Section 6.6 Error Codes -> 400 INVALID_QUERY, 401 INVALID_API_KEY, 403 ACCESS_DENIED, 404 NOT_FOUND",
              "[PASS] 100% Test Pass across all Phase 2 requirements."])

    prs.save(output_path)
    print(f"Phase 2 Presentation saved to: {output_path}")

if __name__ == "__main__":
    out_dir = r"c:\Users\Devil\OneDrive\Desktop\capstone\HarshithKumar_MDDS_Submission\PPT_Slides"
    out_file = os.path.join(out_dir, "Phase_2_API_and_Dashboard_Specification.pptx")
    build_phase2_presentation(out_file)
